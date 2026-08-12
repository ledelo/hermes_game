#!/usr/bin/env python3
"""검수 스크립트 — 「깔끔한 국문 이력서·자소서 세트」

실행: python3 verify.py

검사 항목
  1. 파일        4종 존재·크기
  2. 구조        python-docx 재개봉 → 문단 수, 표 수, 표별 행×열, 병합 셀
  3. 한글        본문에 한글이 실제로 들어 있는지 (인코딩·글리프 유실 검사)
  4. 글꼴        document.xml 안의 rFonts 실측 — 맑은 고딕 외 글꼴 혼입 검사
  5. 색상        w:color 실측 — 포인트 컬러가 1색(남색)인지
  6. 개인정보    주민등록번호·가족사항 등 최신 채용 관행에서 빠진 항목 0건
  7. 금지 문구   합격 보장류·과장 표현 0건 (docx 4종 + README.md)
  8. 렌더        soffice로 PDF 변환 → 쪽수 실측

추정 없이 실측만 출력합니다. 판정은 마지막 줄에 요약합니다.
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

try:
    from docx import Document
except ImportError:
    print("python-docx 가 필요합니다:  pip install python-docx")
    sys.exit(2)

HERE = os.path.dirname(os.path.abspath(__file__))
FILES = [
    "이력서_신입형.docx",
    "이력서_경력형.docx",
    "자기소개서_템플릿.docx",
    "사용가이드.docx",
]
EXPECTED_PAGES = {  # soffice 렌더 기준 상한
    "이력서_신입형.docx": (1, 1),
    "이력서_경력형.docx": (1, 2),
    "자기소개서_템플릿.docx": (1, 2),
    "사용가이드.docx": (1, 1),
}

HANGUL = re.compile(r"[가-힣]")
FONT_OK = {"맑은 고딕"}
COLOR_ALLOWED = {"1F3864", "808080", "1A1A1A", "auto"}  # 남색 포인트 / 안내 회색 / 본문

# 6. 개인정보 최소화 — 서식에 칸이 있으면 안 되는 항목
PRIVACY_TERMS = ["주민등록번호", "주민번호", "가족사항", "가족관계", "본적", "신장", "체중", "혈액형", "종교", "혼인"]
# 아래 문맥에서 발견된 단어는 '넣지 않았다'는 안내 문장이므로 항목이 아닙니다.
PRIVACY_ALLOWED = [r"넣지\s*않았습니다", r"빠지는\s*항목"]

# 7. 금지 문구 — COPY_PACK.md 0.2 공통 규칙 + 채용 상품 특유의 과장 표현
FORBIDDEN = [
    "합격 보장", "합격보장", "합격을 보장", "취업 보장", "취업보장",
    "100%", "무조건", "최저가", "보장", "완전무결", "확실히 붙", "반드시 합격",
    "서류 통과율", "합격률", "합격 확률", "면접 확정",
]
FORBIDDEN_ALLOWED = []  # 예외 없음 — 4종·README 어디에도 쓰지 않는다

ok = True
def head(t):
    print("\n" + t)
    print("-" * 74)

def fail(msg):
    global ok
    ok = False
    print("  [실패] " + msg)


# ── 1. 파일 ──────────────────────────────────────────────────────────────
head("1. 파일")
paths = {}
for f in FILES:
    p = os.path.join(HERE, f)
    if not os.path.exists(p):
        fail(f"{f} 없음")
        continue
    paths[f] = p
    print(f"  {f:26s} {os.path.getsize(p):>7,} bytes")

# ── 2·3. 구조 + 한글 ─────────────────────────────────────────────────────
head("2. 구조 (python-docx 재개봉) / 3. 한글")
for f, p in paths.items():
    doc = Document(p)
    body_paras = len(doc.paragraphs)
    body_text = [x.text for x in doc.paragraphs]
    cell_paras = 0
    tbl_shape = []
    for t in doc.tables:
        rows = len(t.rows)
        cols = len(t.columns)
        tbl_shape.append(f"{rows}×{cols}")
        for row in t.rows:
            for c in row.cells:
                cell_paras += len(c.paragraphs)
                body_text.extend(x.text for x in c.paragraphs)
    text = "\n".join(body_text)
    n_hangul = len(HANGUL.findall(text))
    # 세로 병합(사진 칸) 실측
    with zipfile.ZipFile(p) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    vmerge = xml.count("<w:vMerge")
    print(f"  {f}")
    print(f"     본문 문단 {body_paras:3d} · 표 안 문단 {cell_paras:3d} · 표 {len(doc.tables)}개 [{', '.join(tbl_shape)}]"
          f" · 세로병합 {vmerge} · 한글 {n_hangul:,}자")
    if n_hangul < 100:
        fail(f"{f}: 한글 글자 수 {n_hangul} — 본문이 비었거나 인코딩 문제")
    if len(doc.tables) == 0:
        fail(f"{f}: 표가 없음")

# ── 4. 글꼴 ──────────────────────────────────────────────────────────────
head("4. 글꼴 (rFonts 실측)")
for f, p in paths.items():
    with zipfile.ZipFile(p) as z:
        parts = {n: z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xml")}
    fonts = set()
    for name in ("word/document.xml", "word/styles.xml", "word/numbering.xml"):
        if name in parts:
            for m in re.finditer(r'<w:rFonts([^/>]*)/?>', parts[name]):
                for attr in re.finditer(r'w:(?:ascii|hAnsi|eastAsia|cs)="([^"]+)"', m.group(1)):
                    fonts.add(attr.group(1))
    extra = fonts - FONT_OK
    print(f"  {f:26s} {sorted(fonts)}")
    if extra:
        fail(f"{f}: 맑은 고딕 외 글꼴 혼입 {sorted(extra)}")

# ── 5. 색상 ──────────────────────────────────────────────────────────────
head("5. 색상 (w:color 실측 — 포인트 1색)")
for f, p in paths.items():
    with zipfile.ZipFile(p) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    colors = set(re.findall(r'<w:color w:val="([0-9A-Fa-f]{6}|auto)"', xml))
    extra = {c.upper() for c in colors} - COLOR_ALLOWED
    print(f"  {f:26s} {sorted(colors)}")
    if extra:
        fail(f"{f}: 허용 외 글자색 {sorted(extra)}")

# ── 6·7. 개인정보 항목 / 금지 문구 ───────────────────────────────────────
def all_text(path):
    doc = Document(path)
    out = [x.text for x in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out.extend(x.text for x in c.paragraphs)
    return out

targets = {f: all_text(p) for f, p in paths.items()}
readme = os.path.join(HERE, "README.md")
if os.path.exists(readme):
    with open(readme, encoding="utf-8") as fh:
        targets["README.md"] = fh.read().splitlines()

head("6. 개인정보 항목 (주민등록번호·가족사항 등)")
hits6 = 0
for name, lines in targets.items():
    for i, line in enumerate(lines, 1):
        for term in PRIVACY_TERMS:
            if term in line:
                if any(re.search(pat, line) for pat in PRIVACY_ALLOWED):
                    print(f"  (안내 문장) {name} L{i}: …{line.strip()[:60]}…")
                    continue
                hits6 += 1
                fail(f"{name} L{i}: '{term}' — {line.strip()[:60]}")
print(f"  항목으로 쓰인 사례: {hits6}건")

head("7. 금지 문구")
hits7 = 0
for name, lines in targets.items():
    for i, line in enumerate(lines, 1):
        for term in FORBIDDEN:
            if term in line:
                if any(re.search(pat, line) for pat in FORBIDDEN_ALLOWED):
                    continue
                hits7 += 1
                fail(f"{name} L{i}: '{term}' — {line.strip()[:60]}")
print(f"  검사어 {len(FORBIDDEN)}개 / 대상 {len(targets)}개 파일 / 적발 {hits7}건")

# ── 8. 렌더 ──────────────────────────────────────────────────────────────
head("8. 렌더 (soffice → PDF 쪽수)")
if shutil.which("soffice") is None:
    print("  soffice 없음 — 건너뜀 (apt-get install -y libreoffice-writer)")
else:
    tmp = tempfile.mkdtemp(prefix="resume-kit-verify-")
    for f, p in paths.items():
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, p],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        pdf = os.path.join(tmp, f.replace(".docx", ".pdf"))
        if not os.path.exists(pdf):
            fail(f"{f}: PDF 변환 실패")
            continue
        info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
        m = re.search(r"Pages:\s+(\d+)", info)
        n = int(m.group(1)) if m else -1
        lo, hi = EXPECTED_PAGES[f]
        mark = "" if lo <= n <= hi else f"  ← 기대 {lo}~{hi}쪽"
        print(f"  {f:26s} {n}쪽{mark}")
        if not lo <= n <= hi:
            fail(f"{f}: {n}쪽 (기대 {lo}~{hi}쪽)")
    shutil.rmtree(tmp, ignore_errors=True)

print("\n" + "=" * 74)
print("검수 결과: " + ("통과" if ok else "실패 — 위 [실패] 줄 확인"))
sys.exit(0 if ok else 1)
